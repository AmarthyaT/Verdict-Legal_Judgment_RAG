import os

pptx_path = r"c:\Users\tamar\Desktop\programming\github_repos\rag_proto\verdict_Amarthya_1AY23CD066.pptx"
output_txt_path = r"C:\Users\tamar\.gemini\antigravity\brain\bdabaa9c-3007-44d8-bcff-9330bb54c764\scratch\pptx_text.txt"

try:
    from pptx import Presentation
    print("python-pptx is installed. Extracting slides...")
    prs = Presentation(pptx_path)
    text = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        slide_text.append(f"=== SLIDE {i+1} ===")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        text.append("\n".join(slide_text))
    
    with open(output_txt_path, 'w', encoding='utf-8') as f:
        f.write("\n\n".join(text))
    print(f"Extraction successful. Text saved to {output_txt_path}")
except ImportError:
    print("python-pptx is not installed. We will skip slide extraction and rely on the research paper.")
except Exception as e:
    print(f"Error: {e}")
